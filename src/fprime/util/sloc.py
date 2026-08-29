"""fprime.util.sloc: SLOC counting for F Prime projects

Implements the 'fprime-util sloc' command. Segments an F Prime project into modules (any
directory whose CMakeLists.txt calls register_fprime_*, add_library, or add_executable), counts
SLOC per module using pygount, and rolls totals up per section (project, framework, libraries)
and for the whole repository.

The command is context aware:
    - In a component/module directory: counts just that module
    - In a deployment directory: counts the deployment and its dependencies (needs build cache)
    - At the project root: counts everything, segmented by project/framework/libraries
    - With --recursive: counts all modules under the current directory

FPP files are supported through the FPP pygments lexer plugin (fprime.fpp.lexer) which pygount
picks up automatically.

@author lestarch
"""

import argparse
import json
import re
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

from pygount import SourceAnalysis, SourceState

# openpyxl is optional: only needed for --excel-report
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:
    Workbook = None
    Font = None

from fprime.fbuild.builder import Build

MODULE_PATTERN = re.compile(
    r"^\s*(register_fprime_\w+|add_library|add_executable)\s*\(", re.MULTILINE
)
DEPLOYMENT_PATTERN = re.compile(r"^\s*register_fprime_deployment\s*\(", re.MULTILINE)
LIBRARY_TOKEN_PATTERN = re.compile(r"lib([A-Za-z0-9_\-]+)\.(?:a|so|dylib)")
SUBMODULE_PATH_PATTERN = re.compile(r"^\s*path\s*=\s*(.+)$", re.MULTILINE)

SKIP_DIRECTORY_NAMES = {"__pycache__", "node_modules", "build-artifacts"}
NON_CODE_LANGUAGES = {
    "Text only",
    "Markdown",
    "reStructuredText",
    "__unknown__",
    "__binary__",
    "__empty__",
    "__error__",
}
UT_PATH_NAMES = {"test", "tests", "ut"}
OTHER_MODULE_NAME = "(non-module files)"

# Language classification: FPP model files, CMake, C/C++ (with UT and autocode splits), other
CPP_LANGUAGES = {"C++", "C"}
CATEGORY_FPP = "FPP"
CATEGORY_CMAKE = "CMake"
CATEGORY_CPP = "C/C++"
CATEGORY_CPP_AC = "C/C++ AC"
CATEGORY_CPP_UT = "C/C++ UT"
CATEGORY_CPP_AC_UT = "C/C++ AC UT"
BASE_CATEGORIES = [
    CATEGORY_FPP,
    CATEGORY_CMAKE,
    CATEGORY_CPP,
    CATEGORY_CPP_AC,
    CATEGORY_CPP_UT,
    CATEGORY_CPP_AC_UT,
]
AC_CATEGORIES = {CATEGORY_CPP_AC, CATEGORY_CPP_AC_UT}
AC_UT_FILE_PATTERN = re.compile(r"(TesterBase|GTestBase|TesterHelpers)")


@dataclass
class Counts:
    """Line counts for a set of files"""

    files: int = 0
    code: int = 0
    comment: int = 0
    blank: int = 0

    def accumulate(self, other: "Counts"):
        """Add another set of counts into this one"""
        self.files += other.files
        self.code += other.code
        self.comment += other.comment
        self.blank += other.blank

    def to_dict(self) -> dict:
        """Dictionary representation for JSON output"""
        return {
            "files": self.files,
            "code": self.code,
            "comment": self.comment,
            "blank": self.blank,
        }


def classify(language: str, is_ut: bool, is_ac: bool) -> str:
    """Map a pygount language (and UT/AC context) to a report category"""
    if language == "FPP":
        return CATEGORY_FPP
    if language == "CMake":
        return CATEGORY_CMAKE
    if language in CPP_LANGUAGES:
        if is_ac:
            return CATEGORY_CPP_AC_UT if is_ut else CATEGORY_CPP_AC
        return CATEGORY_CPP_UT if is_ut else CATEGORY_CPP
    return language


@dataclass
class ModuleSloc:
    """SLOC results for a single module, split into categories"""

    name: str
    path: Path
    section: str
    is_deployment: bool = False
    categories: Dict[str, Counts] = field(default_factory=dict)

    def add(self, category: str, counts: Counts):
        """Accumulate counts into a category"""
        self.categories.setdefault(category, Counts()).accumulate(counts)

    def category(self, category: str) -> Counts:
        """Counts for a category (zero counts if absent)"""
        return self.categories.get(category, Counts())

    def cpp_total(self) -> int:
        """Total C/C++ code lines across code, UT, and autocode"""
        return sum(
            self.category(category).code
            for category in (
                CATEGORY_CPP,
                CATEGORY_CPP_AC,
                CATEGORY_CPP_UT,
                CATEGORY_CPP_AC_UT,
            )
        )

    def other_languages(self) -> List[str]:
        """Languages outside the base categories"""
        return [c for c in self.categories if c not in BASE_CATEGORIES]

    def other_total(self) -> int:
        """Total code lines across all 'other' languages"""
        return sum(self.categories[c].code for c in self.other_languages())

    def total_code(self) -> int:
        """Total code lines across all categories"""
        return sum(counts.code for counts in self.categories.values())

    def total_files(self) -> int:
        """Total files across all categories"""
        return sum(counts.files for counts in self.categories.values())


@dataclass
class SlocReport:
    """Full report: modules grouped into sections"""

    modules: List[ModuleSloc] = field(default_factory=list)
    languages: Dict[str, Counts] = field(default_factory=dict)
    ac_counted: bool = False

    def sections(self) -> List[str]:
        """Ordered unique section names"""
        seen = {}
        for module in self.modules:
            seen.setdefault(module.section, None)
        return list(seen.keys())

    def other_languages(self) -> List[str]:
        """All 'other' languages in the report, largest first"""
        totals: Dict[str, int] = {}
        for module in self.modules:
            for language in module.other_languages():
                totals[language] = (
                    totals.get(language, 0) + module.categories[language].code
                )
        return sorted(totals, key=lambda language: -totals[language])

    def section_totals(self, section: str) -> ModuleSloc:
        """Totals for a section, as a pseudo-module"""
        return self._totals([m for m in self.modules if m.section == section])

    def grand_totals(self) -> ModuleSloc:
        """Totals for the whole report, as a pseudo-module"""
        return self._totals(self.modules)

    @staticmethod
    def _totals(modules: List[ModuleSloc]) -> ModuleSloc:
        total = ModuleSloc(name="TOTAL", path=Path("."), section="")
        for module in modules:
            for category, counts in module.categories.items():
                total.add(category, counts)
        return total


def is_build_cache(path: Path) -> bool:
    """Check if a directory is (or holds) a build cache"""
    return (
        (path / "CMakeCache.txt").exists()
        or (path / ".fprime-build-dir").exists()
        or path.name.startswith("build-fprime")
    )


def should_skip_directory(path: Path) -> bool:
    """Check if a directory should be excluded from source scanning"""
    return (
        path.name.startswith(".")
        or path.name in SKIP_DIRECTORY_NAMES
        or is_build_cache(path)
        or (path / "pyvenv.cfg").exists()  # virtual environments
    )


def read_submodule_paths(root: Path) -> List[Path]:
    """Read git submodule paths under a root, if any"""
    gitmodules = root / ".gitmodules"
    if not gitmodules.is_file():
        return []
    matches = SUBMODULE_PATH_PATTERN.findall(gitmodules.read_text(errors="replace"))
    return [(root / match.strip()).resolve() for match in matches]


def registers_module(cmake_lists: Path) -> Tuple[bool, bool]:
    """Check a CMakeLists.txt: returns (is_module, is_deployment)"""
    try:
        content = cmake_lists.read_text(errors="replace")
    except OSError:
        return False, False
    return (
        MODULE_PATTERN.search(content) is not None,
        DEPLOYMENT_PATTERN.search(content) is not None,
    )


def discover_modules(
    root: Path, section: str, excludes: List[Path]
) -> Tuple[List[ModuleSloc], Dict[Path, ModuleSloc]]:
    """Discover modules under a root and attribute every source file to one

    Files not under any module directory are attributed to a per-section pseudo-module. Returns
    the module list and a mapping of file path to owning module.
    """
    modules: Dict[Path, ModuleSloc] = {}
    file_owners: Dict[Path, ModuleSloc] = {}
    other = ModuleSloc(name=OTHER_MODULE_NAME, path=root, section=section)

    def walk(directory: Path, owner: Optional[ModuleSloc]):
        cmake_lists = directory / "CMakeLists.txt"
        if cmake_lists.is_file():
            is_module, is_deployment = registers_module(cmake_lists)
            if is_module:
                name = (
                    str(directory.relative_to(root)).replace("/", "_")
                    if directory != root
                    else root.name
                )
                owner = modules.setdefault(
                    directory,
                    ModuleSloc(
                        name=name,
                        path=directory,
                        section=section,
                        is_deployment=is_deployment,
                    ),
                )
        for entry in sorted(directory.iterdir()):
            if entry.is_symlink():
                continue
            if entry.is_dir():
                if should_skip_directory(entry) or entry.resolve() in excludes:
                    continue
                walk(entry, owner)
            elif entry.is_file():
                file_owners[entry] = owner if owner is not None else other
        return

    walk(root, None)
    module_list = list(modules.values())
    if any(owner is other for owner in file_owners.values()):
        module_list.append(other)
    return module_list, file_owners


def is_ut_file(path: Path, module_path: Path) -> bool:
    """Check if a file is unit-test code based on its path within the module"""
    try:
        relative = path.relative_to(module_path)
    except ValueError:
        return False
    return any(part.lower() in UT_PATH_NAMES for part in relative.parts[:-1])


def analyze_file(path: Path) -> Optional[SourceAnalysis]:
    """Run pygount on a single file, returning None for non-code files"""
    try:
        # pygount treats any *.txt (including CMakeLists.txt) as plain text; reanalyze
        # CMakeLists.txt under a .cmake name so it counts as CMake
        if path.name == "CMakeLists.txt":
            with tempfile.NamedTemporaryFile(
                "wb", suffix=".cmake", delete=False
            ) as temporary:
                temporary_path = Path(temporary.name)
            try:
                temporary_path.write_bytes(path.read_bytes())
                return analyze_file(temporary_path)
            finally:
                temporary_path.unlink()
        analysis = SourceAnalysis.from_file(
            str(path), "sloc", fallback_encoding="utf-8"
        )
    except Exception:  # pygount can raise on odd encodings; skip such files
        return None
    if (
        analysis.state != SourceState.analyzed
        or analysis.language in NON_CODE_LANGUAGES
    ):
        return None
    return analysis


def count_files(
    file_owners: Dict[Path, ModuleSloc], languages: Dict[str, Counts]
) -> None:
    """Count all files, accumulating into their owning modules and the language rollup"""
    for path, module in file_owners.items():
        analysis = analyze_file(path)
        if analysis is None:
            continue
        counts = Counts(
            files=1,
            code=analysis.code_count + analysis.string_count,
            comment=analysis.documentation_count,
            blank=analysis.empty_count,
        )
        module.add(
            classify(analysis.language, is_ut_file(path, module.path), is_ac=False),
            counts,
        )
        languages.setdefault(analysis.language, Counts()).accumulate(counts)


def count_autocoded(
    build: Build, modules: List[ModuleSloc], languages: Dict[str, Counts]
) -> bool:
    """Count autocoded files in the build cache for each module

    Returns True if the build cache was available and AC files were counted.
    """
    if build.build_dir is None or not (build.build_dir / "CMakeCache.txt").exists():
        return False
    for module in modules:
        try:
            cache_path = build.get_build_cache_path(module.path)
        except Exception:
            continue
        # Generated files land directly in the module's mirrored cache directory;
        # do not recurse as subdirectories belong to other modules or CMake internals
        for path in sorted(cache_path.iterdir()):
            if not path.is_file():
                continue
            if path.suffix not in {".hpp", ".cpp", ".h", ".c", ".fppi", ".py"}:
                continue
            analysis = analyze_file(path)
            if analysis is None:
                continue
            counts = Counts(
                files=1,
                code=analysis.code_count + analysis.string_count,
                comment=analysis.documentation_count,
                blank=analysis.empty_count,
            )
            is_ut = AC_UT_FILE_PATTERN.search(path.name) is not None
            module.add(classify(analysis.language, is_ut, is_ac=True), counts)
            languages.setdefault(analysis.language, Counts()).accumulate(counts)
    return True


def get_sections(
    build: Build, include_submodules: bool
) -> List[Tuple[str, Path, List[Path]]]:
    """Compute the (name, root, excludes) sections: project, framework, and libraries"""
    framework_setting = build.get_settings("framework_path", None)
    if framework_setting is None:
        raise SlocException(
            "Cannot determine 'framework_path'. Run from a directory with a settings.ini or "
            "supply an existing build cache with --build-cache."
        )
    framework = Path(framework_setting).resolve()
    project = Path(build.get_settings("project_root", framework)).resolve()
    libraries = [
        Path(library).resolve()
        for library in build.get_settings("library_locations", [])
    ]
    sections = []
    roots = []
    if project != framework and project not in libraries:
        sections.append((f"project: {project.name}", project))
    sections.append(("fprime", framework))
    sections.extend((f"library: {library.name}", library) for library in libraries)
    roots = [root for _, root in sections]
    results = []
    for name, root in sections:
        excludes = [
            other for other in roots if other != root and other.is_relative_to(root)
        ]
        if not include_submodules:
            excludes.extend(read_submodule_paths(root))
        results.append((name, root, excludes))
    return results


def find_section_name(path: Path, sections: List[Tuple[str, Path, List[Path]]]) -> str:
    """Find the section name owning a path, preferring the deepest matching root"""
    best = None
    for name, root, _ in sections:
        if path == root or path.is_relative_to(root):
            if best is None or len(root.parts) > len(best[1].parts):
                best = (name, root)
    return best[0] if best else "project"


def ninja_link_tokens(ninja_file: Path, executable: str) -> set:
    """Extract library names linked into an executable from a build.ninja file

    Walks the ninja build graph from the build statement(s) producing the executable, following
    dependencies transitively, and collects the lib<name> tokens seen along the way.
    """
    text = ninja_file.read_text(errors="replace").replace("$\n", " ")
    edges: Dict[str, List[str]] = {}
    for line in text.splitlines():
        if not line.startswith("build ") or ":" not in line:
            continue
        outputs, _, remainder = line[len("build ") :].partition(":")
        dependencies = remainder.replace("|", " ").split()[1:]  # drop the rule name
        for output in outputs.replace("|", " ").split():
            edges.setdefault(output, []).extend(dependencies)
    # deployment targets may be mangled with path prefixes (e.g. Project_Deployment)
    starts = [
        output
        for output in edges
        if Path(output).name == executable
        or Path(output).name.endswith(f"_{executable}")
    ]
    tokens = set()
    visited = set()
    queue = list(starts)
    while queue:
        node = queue.pop()
        if node in visited:
            continue
        visited.add(node)
        tokens.update(LIBRARY_TOKEN_PATTERN.findall(node))
        queue.extend(edges.get(node, []))
    return tokens


def deployment_module_paths(
    build: Build, deployment: Path, modules: List[ModuleSloc]
) -> List[ModuleSloc]:
    """Restrict modules to a deployment and its linked dependencies

    Parses link lines (link.txt for Makefiles, build.ninja for Ninja) in the build cache to find
    the libraries linked into the deployment, then matches them to discovered module names.
    """
    tokens = set()
    try:
        cache_path = build.get_build_cache_path(deployment)
        for link_file in cache_path.rglob("link.txt"):
            tokens.update(
                LIBRARY_TOKEN_PATTERN.findall(link_file.read_text(errors="replace"))
            )
    except Exception:
        pass
    ninja_file = build.build_dir / "build.ninja" if build.build_dir else None
    if not tokens and ninja_file is not None and ninja_file.exists():
        tokens.update(ninja_link_tokens(ninja_file, deployment.name))
    if not tokens:
        raise SlocException(
            f"Could not determine dependencies for deployment '{deployment.name}' from the "
            "build cache. Ensure the deployment has been generated and built."
        )
    selected = []
    for module in modules:
        in_deployment = module.path == deployment or module.path.is_relative_to(
            deployment
        )
        if in_deployment or module.name in tokens:
            selected.append(module)
    return selected


def discover_all(
    scope_roots: List[Tuple[str, Path, List[Path]]],
) -> Tuple[List[ModuleSloc], Dict[Path, ModuleSloc]]:
    """Discover modules and file ownership under each scope root"""
    all_modules: List[ModuleSloc] = []
    all_owners: Dict[Path, ModuleSloc] = {}
    for name, root, excludes in scope_roots:
        modules, file_owners = discover_modules(root, name, excludes)
        all_modules.extend(modules)
        all_owners.update(file_owners)
    return all_modules, all_owners


def build_report(
    build: Build,
    modules: List[ModuleSloc],
    file_owners: Dict[Path, ModuleSloc],
    with_autocode: bool,
) -> SlocReport:
    """Count SLOC for the given modules and produce a report"""
    report = SlocReport(modules=list(modules))
    selected = set(id(module) for module in modules)
    owners = {
        path: module for path, module in file_owners.items() if id(module) in selected
    }
    count_files(owners, report.languages)
    if with_autocode:
        report.ac_counted = count_autocoded(build, report.modules, report.languages)
    report.modules.sort(key=lambda module: (module.section, module.name))
    return report


####
# Output rendering
####
def report_columns(report: SlocReport) -> List[str]:
    """Column headers: base categories, C/C++ total, per-language others, other/grand totals"""
    columns = ["Module", "Files"] + BASE_CATEGORIES + ["C/C++ Total"]
    columns += [f"other ({language})" for language in report.other_languages()]
    return columns + ["other (Total)", "Total"]


def module_row(module: ModuleSloc, report: SlocReport, label: str = None) -> List[str]:
    """Build a display row for a module (or totals pseudo-module)"""

    def cell(category: str) -> str:
        if category in AC_CATEGORIES and not report.ac_counted:
            return "-"
        return str(module.category(category).code)

    row = [label if label is not None else module.name, str(module.total_files())]
    row += [cell(category) for category in BASE_CATEGORIES]
    row.append(str(module.cpp_total()))
    row += [
        str(module.category(language).code) for language in report.other_languages()
    ]
    return row + [str(module.other_total()), str(module.total_code())]


def render_table(report: SlocReport, output=None):
    """Render the report as an aligned text table to the given stream"""
    output = output if output is not None else sys.stdout
    columns = report_columns(report)
    rows: List[List[str]] = []
    separators: List[int] = []
    for section in report.sections():
        rows.append([f"[{section}]"] + [""] * (len(columns) - 1))
        rows.extend(
            module_row(module, report)
            for module in report.modules
            if module.section == section
        )
        rows.append(
            module_row(report.section_totals(section), report, f"TOTAL [{section}]")
        )
        separators.append(len(rows))
    if len(report.sections()) > 1:
        rows.append(module_row(report.grand_totals(), report, "GRAND TOTAL"))
    widths = [
        max(len(columns[i]), max((len(row[i]) for row in rows), default=0))
        for i in range(len(columns))
    ]

    def print_row(cells):
        padded = [cells[0].ljust(widths[0])] + [
            cell.rjust(width) for cell, width in zip(cells[1:], widths[1:])
        ]
        print("  ".join(padded), file=output)

    print_row(columns)
    print("-" * (sum(widths) + 2 * (len(widths) - 1)), file=output)
    for index, row in enumerate(rows, 1):
        print_row(row)
        if index in separators:
            print("-" * (sum(widths) + 2 * (len(widths) - 1)), file=output)


def render_markdown(report: SlocReport) -> str:
    """Render the report as a Markdown document"""
    columns = report_columns(report)
    lines = ["# SLOC Report", ""]
    if report.ac_counted:
        lines += ["Autocoded (AC) files counted from the build cache.", ""]
    for section in report.sections():
        lines += [f"## {section}", "", "| " + " | ".join(columns) + " |"]
        lines.append("|" + "|".join(["---"] * len(columns)) + "|")
        lines.extend(
            "| " + " | ".join(module_row(module, report)) + " |"
            for module in report.modules
            if module.section == section
        )
        lines.append(
            "| "
            + " | ".join(
                f"**{cell}**" if cell else ""
                for cell in module_row(report.section_totals(section), report, "TOTAL")
            )
            + " |"
        )
        lines.append("")
    if len(report.sections()) > 1:
        lines += ["## Grand Total", "", "| " + " | ".join(columns) + " |"]
        lines.append("|" + "|".join(["---"] * len(columns)) + "|")
        lines.append(
            "| "
            + " | ".join(
                f"**{cell}**" if cell else ""
                for cell in module_row(report.grand_totals(), report, "ALL")
            )
            + " |"
        )
        lines.append("")
    return "\n".join(lines) + "\n"


def render_excel(report: SlocReport, path: Path):
    """Render the report as an Excel workbook"""
    if Workbook is None:
        raise SlocException(
            "openpyxl is not installed. Install it (e.g. 'pip install openpyxl') to use --excel-report."
        )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SLOC"
    bold = Font(bold=True)

    def append_row(cells: List[str], is_bold: bool = False):
        sheet.append(
            [int(cell) if cell.lstrip("-").isdigit() else cell for cell in cells]
        )
        if is_bold:
            for cell in sheet[sheet.max_row]:
                cell.font = bold

    append_row(report_columns(report), is_bold=True)
    for section in report.sections():
        append_row([f"[{section}]"], is_bold=True)
        for module in report.modules:
            if module.section == section:
                append_row(module_row(module, report))
        append_row(
            module_row(report.section_totals(section), report, f"TOTAL [{section}]"),
            is_bold=True,
        )
    if len(report.sections()) > 1:
        append_row(
            module_row(report.grand_totals(), report, "GRAND TOTAL"), is_bold=True
        )
    sheet.column_dimensions["A"].width = max(
        len(str(row[0].value or "")) for row in sheet.iter_rows(min_col=1, max_col=1)
    )
    sheet.freeze_panes = "B2"
    workbook.save(path)


def categories_dict(module: ModuleSloc) -> dict:
    """Per-category counts dictionary for JSON output"""
    return {
        category: counts.to_dict() for category, counts in module.categories.items()
    }


def render_json(report: SlocReport) -> str:
    """Render the report as JSON"""
    sections = {}
    for section in report.sections():
        sections[section] = {
            "modules": {
                module.name: {
                    "path": str(module.path),
                    "deployment": module.is_deployment,
                    "categories": categories_dict(module),
                }
                for module in report.modules
                if module.section == section
            },
            "totals": categories_dict(report.section_totals(section)),
        }
    return json.dumps(
        {
            "autocode_counted": report.ac_counted,
            "sections": sections,
            "totals": categories_dict(report.grand_totals()),
            "languages": {
                language: counts.to_dict()
                for language, counts in report.languages.items()
            },
        },
        indent=2,
    )


####
# CLI handling
####
def find_context_module(path: Path) -> Optional[Tuple[Path, bool]]:
    """Find the nearest ancestor (inclusive) module directory of a path

    Returns (module_directory, is_deployment) or None when no module is found.
    """
    for candidate in [path, *path.parents]:
        cmake_lists = candidate / "CMakeLists.txt"
        if cmake_lists.is_file():
            is_module, is_deployment = registers_module(cmake_lists)
            if is_module:
                return candidate, is_deployment
    return None


def run_sloc(
    build: Build,
    parsed: argparse.Namespace,
    _: Dict[str, str],
    __: Dict[str, str],
    ___,
):
    """Run the sloc command: context-aware SLOC counting

    Args:
        build: build object (cache validation skipped; used for settings and AC counting)
        parsed: parsed command namespace
        _: unused cmake arguments
        __: unused make arguments
        ___: unused pass through arguments
    """
    path = Path(parsed.path).resolve()
    sections = get_sections(build, parsed.include_submodules)
    section_roots = [root for _, root, _ in sections]

    def scoped_excludes(scope: Path) -> List[Path]:
        """Section excludes (including submodules) that fall under a scope directory"""
        return [
            exclude
            for _, _, excludes in sections
            for exclude in excludes
            if exclude.is_relative_to(scope)
        ]

    cache_available = (
        build.build_dir is not None and (build.build_dir / "CMakeCache.txt").exists()
    )

    if parsed.recursive:
        if not cache_available:
            raise SlocException(
                "--recursive requires a build cache. Run 'fprime-util generate' first."
            )
        name = find_section_name(path, sections)
        modules, file_owners = discover_all([(name, path, scoped_excludes(path))])
        report = build_report(build, modules, file_owners, with_autocode=True)
    elif (
        path in section_roots
        or path == Path(build.get_settings("project_root", path)).resolve()
    ):
        modules, file_owners = discover_all(sections)
        report = build_report(
            build, modules, file_owners, with_autocode=cache_available
        )
    else:
        context = find_context_module(path)
        if context is None:
            raise SlocException(
                f"'{path}' is not part of an F Prime module. Use --recursive to count a "
                "directory tree."
            )
        module_path, is_deployment = context
        name = find_section_name(module_path, sections)
        if is_deployment:
            if not cache_available:
                raise SlocException(
                    "Deployment SLOC counts require a build cache. Run 'fprime-util generate' "
                    "first."
                )
            # Discover everywhere, restrict to the deployment closure, then count
            modules, file_owners = discover_all(sections)
            selected = deployment_module_paths(build, module_path, modules)
            report = build_report(build, selected, file_owners, with_autocode=True)
        else:
            modules, file_owners = discover_all(
                [(name, module_path, scoped_excludes(module_path))]
            )
            report = build_report(
                build, modules, file_owners, with_autocode=cache_available
            )

    render_table(report)
    if parsed.markdown_report:
        Path(parsed.markdown_report).write_text(render_markdown(report))
        print(f"[INFO] Markdown report written to {parsed.markdown_report}")
    if parsed.json_report:
        Path(parsed.json_report).write_text(render_json(report))
        print(f"[INFO] JSON report written to {parsed.json_report}")
    if parsed.excel_report:
        render_excel(report, Path(parsed.excel_report))
        print(f"[INFO] Excel report written to {parsed.excel_report}")
    return 0


def add_sloc_parsers(
    subparsers, common: argparse.ArgumentParser, help_text
) -> Tuple[Dict[str, Callable], Dict[str, argparse.ArgumentParser]]:
    """Add the sloc subparser to the CLI

    Args:
        subparsers: subparsers to add the sloc parser to
        common: common parent parser
        help_text: HelpText class providing help messages

    Returns:
        Tuple of runner dictionary and parser dictionary
    """
    sloc_parser = subparsers.add_parser(
        "sloc",
        description=help_text.long("sloc"),
        help=help_text.short("sloc"),
        parents=[common],
        add_help=False,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    sloc_parser.add_argument(
        "--recursive",
        default=False,
        action="store_true",
        help="Count all modules under the current directory (requires a build cache)",
    )
    sloc_parser.add_argument(
        "--include-submodules",
        default=False,
        action="store_true",
        help="Include git submodules (e.g. googletest) in counts",
    )
    sloc_parser.add_argument(
        "--markdown-report",
        default=None,
        type=Path,
        help="Write a Markdown report to the given path",
    )
    sloc_parser.add_argument(
        "--json-report",
        default=None,
        type=Path,
        help="Write a JSON report to the given path",
    )
    sloc_parser.add_argument(
        "--excel-report",
        default=None,
        type=Path,
        help="Write an Excel (.xlsx) report to the given path",
    )
    return {"sloc": run_sloc}, {"sloc": sloc_parser}


class SlocException(Exception):
    """An exception in SLOC counting"""
